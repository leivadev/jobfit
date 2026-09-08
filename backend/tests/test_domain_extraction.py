import io

import pypdf
import pytest
from docx import Document

from backend.domain.extraction import (
    CvExtractionError,
    UnsupportedCvFormatError,
    extract_text,
)


def make_pdf_bytes(*page_texts: str) -> bytes:
    """Hand-crafted PDF with one content-stream page per `page_texts` entry.

    No PDF-writing library is a project dependency, so tests build the
    minimal valid PDF structure pypdf needs to read it back.
    """
    n = len(page_texts)
    font_id = 3 + 2 * n
    kids = b" ".join(f"{3 + i} 0 R".encode() for i in range(n))
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, n),
    ]
    for i in range(n):
        content_id = 3 + n + i
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R"
                f" /Resources << /Font << /F1 {font_id} 0 R >> >>"
                f" /MediaBox [0 0 612 792] /Contents {content_id} 0 R >>"
            ).encode()
        )
    for text in page_texts:
        content = f"BT /F1 24 Tf 72 712 Td ({text}) Tj ET".encode("latin-1")
        objects.append(
            b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content)
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(buf.tell())
        buf.write(f"{i} 0 obj\n".encode())
        buf.write(obj)
        buf.write(b"\nendobj\n")
    xref_offset = buf.tell()
    total = len(objects) + 1
    buf.write(f"xref\n0 {total}\n".encode())
    buf.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        buf.write(f"{off:010d} 00000 n \n".encode())
    buf.write(b"trailer\n")
    buf.write(f"<< /Size {total} /Root 1 0 R >>\n".encode())
    buf.write(b"startxref\n")
    buf.write(f"{xref_offset}\n".encode())
    buf.write(b"%%EOF")
    return buf.getvalue()


def make_encrypted_pdf_bytes(
    *page_texts: str, user_password: str, owner_password: str | None = None
) -> bytes:
    """Encrypt a hand-crafted PDF with the given user/owner passwords."""
    reader = pypdf.PdfReader(io.BytesIO(make_pdf_bytes(*page_texts)))
    writer = pypdf.PdfWriter(clone_from=reader)
    writer.encrypt(user_password=user_password, owner_password=owner_password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row, cells in zip(table.rows, table_rows):
            for cell, text in zip(row.cells, cells):
                cell.text = text
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extract_text_from_pdf_by_mime_type():
    content = make_pdf_bytes("Jane Doe, Backend Engineer")

    text = extract_text(content, mime_type="application/pdf")

    assert "Jane Doe, Backend Engineer" in text


def test_extract_text_from_pdf_by_extension():
    content = make_pdf_bytes("Jane Doe, Backend Engineer")

    text = extract_text(content, filename="cv.pdf")

    assert "Jane Doe, Backend Engineer" in text


def test_pdf_joins_multiple_pages_with_newline():
    content = make_pdf_bytes("Page One", "Page Two")

    text = extract_text(content, mime_type="application/pdf")

    lines = [line for line in text.splitlines() if line.strip()]
    assert lines == ["Page One", "Page Two"]


def test_pdf_with_owner_password_only_extracts_successfully():
    content = make_encrypted_pdf_bytes(
        "Jane Doe, Backend Engineer", user_password="", owner_password="secret"
    )

    text = extract_text(content, mime_type="application/pdf")

    assert "Jane Doe, Backend Engineer" in text


def test_pdf_with_real_user_password_raises_password_protected_error():
    content = make_encrypted_pdf_bytes(
        "Jane Doe, Backend Engineer", user_password="letmein"
    )

    with pytest.raises(CvExtractionError):
        extract_text(content, mime_type="application/pdf")


def test_extract_text_from_docx_by_mime_type():
    content = make_docx_bytes(["Jane Doe", "Backend Engineer, 5 years"])

    text = extract_text(
        content,
        mime_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
    )

    assert "Jane Doe" in text
    assert "Backend Engineer, 5 years" in text


def test_extract_text_from_docx_by_extension():
    content = make_docx_bytes(["Jane Doe"])

    text = extract_text(content, filename="cv.docx")

    assert "Jane Doe" in text


def test_extract_text_from_docx_includes_table_content():
    content = make_docx_bytes(
        ["Jane Doe"],
        table_rows=[["Skill", "Level"], ["Python", "Expert"]],
    )

    text = extract_text(content, filename="cv.docx")

    assert "Jane Doe" in text
    assert "Python" in text
    assert "Expert" in text


def test_extract_text_from_plain_text_by_mime_type():
    content = b"Jane Doe\nBackend Engineer"

    text = extract_text(content, mime_type="text/plain")

    assert text == "Jane Doe\nBackend Engineer"


def test_extract_text_from_plain_text_by_extension():
    content = b"Jane Doe\nBackend Engineer"

    text = extract_text(content, filename="cv.txt")

    assert text == "Jane Doe\nBackend Engineer"


def test_plain_text_raises_clear_error_for_invalid_utf8():
    content = "Jané Döe".encode("latin-1")

    with pytest.raises(CvExtractionError):
        extract_text(content, mime_type="text/plain")


def test_pdf_raises_extraction_error_for_empty_text():
    content = make_pdf_bytes("")

    with pytest.raises(CvExtractionError):
        extract_text(content, mime_type="application/pdf")


def test_pdf_raises_extraction_error_for_whitespace_only_text():
    content = make_pdf_bytes("   ")

    with pytest.raises(CvExtractionError):
        extract_text(content, mime_type="application/pdf")


def test_docx_raises_extraction_error_for_empty_text():
    content = make_docx_bytes([])

    with pytest.raises(CvExtractionError):
        extract_text(content, filename="cv.docx")


def test_docx_raises_extraction_error_for_whitespace_only_text():
    content = make_docx_bytes(["   ", ""])

    with pytest.raises(CvExtractionError):
        extract_text(content, filename="cv.docx")


def test_plain_text_raises_extraction_error_for_empty_content():
    with pytest.raises(CvExtractionError):
        extract_text(b"", mime_type="text/plain")


def test_plain_text_raises_extraction_error_for_whitespace_only_content():
    with pytest.raises(CvExtractionError):
        extract_text(b"   \n\t  ", mime_type="text/plain")


def test_mime_type_takes_precedence_over_extension():
    content = b"Jane Doe"

    text = extract_text(content, filename="cv.pdf", mime_type="text/plain")

    assert text == "Jane Doe"


def test_falls_back_to_extension_when_mime_type_unrecognized():
    content = b"Jane Doe"

    text = extract_text(
        content, filename="cv.txt", mime_type="application/octet-stream"
    )

    assert text == "Jane Doe"


def test_pdf_raises_extraction_error_for_corrupted_content():
    content = b"not a real pdf"

    with pytest.raises(CvExtractionError):
        extract_text(content, mime_type="application/pdf")


def test_docx_raises_extraction_error_for_corrupted_content():
    content = b"not a real docx"

    with pytest.raises(CvExtractionError):
        extract_text(content, filename="cv.docx")


def test_raises_for_unsupported_format():
    content = b"whatever"

    with pytest.raises(UnsupportedCvFormatError):
        extract_text(content, filename="cv.exe", mime_type="application/x-msdownload")


def test_raises_when_no_filename_or_mime_type():
    with pytest.raises(UnsupportedCvFormatError):
        extract_text(b"whatever")
